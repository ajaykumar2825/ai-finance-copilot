from __future__ import annotations

import csv
import io
import logging
from typing import Any

logger = logging.getLogger(__name__)


def parse_pdf(
    content: bytes,
    filename: str = "",
) -> dict[str, Any]:
    """Extract text from a PDF using pypdf.

    Returns:
        ``{"text": ..., "metadata": {page_count, ...}}``
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    full_text = "\n\n".join(pages)

    # Build page-aware metadata
    page_metadata = []
    for idx, page_text in enumerate(pages):
        page_metadata.append(
            {
                "page_number": idx + 1,
                "text": page_text,
                "start_char": sum(len(p) for p in pages[:idx]),
                "end_char": sum(len(p) for p in pages[: idx + 1]),
            }
        )

    return {
        "text": full_text,
        "metadata": {
            "filename": filename,
            "file_type": "pdf",
            "page_count": len(pages),
            "pages": page_metadata,
        },
    }


def parse_docx(
    content: bytes,
    filename: str = "",
) -> dict[str, Any]:
    """Extract text from a DOCX file using python-docx.

    Returns:
        ``{"text": ..., "metadata": {...}}``
    """
    from docx import Document

    doc = Document(io.BytesIO(content))

    # Combine paragraphs and table text
    parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    full_text = "\n\n".join(parts)

    return {
        "text": full_text,
        "metadata": {
            "filename": filename,
            "file_type": "docx",
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        },
    }


def parse_csv(
    content: bytes,
    filename: str = "",
) -> dict[str, Any]:
    """Parse a CSV file using Python's csv module.

    Returns:
        ``{"text": ..., "metadata": {...}}``
    """
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))

    if not reader.fieldnames:
        return {
            "text": "",
            "metadata": {"filename": filename, "file_type": "csv", "rows": 0},
        }

    headers = reader.fieldnames
    rows = list(reader)
    row_count = len(rows)

    # Build a readable text representation
    parts = [f"Columns: {', '.join(headers)}"]
    for row in rows[:100]:  # cap at 100 rows to avoid huge context
        parts.append(" | ".join(f"{k}={v}" for k, v in row.items() if v))

    full_text = "\n".join(parts)

    return {
        "text": full_text,
        "metadata": {
            "filename": filename,
            "file_type": "csv",
            "rows": row_count,
            "columns": headers,
        },
    }


def parse_xlsx(
    content: bytes,
    filename: str = "",
) -> dict[str, Any]:
    """Parse an Excel spreadsheet using openpyxl (via pandas).

    Returns:
        ``{"text": ..., "metadata": {...}}``
    """
    import pandas as pd

    excel_file = pd.ExcelFile(io.BytesIO(content))
    sheet_names = excel_file.sheet_names
    parts: list[str] = []
    total_rows = 0

    for sheet in sheet_names:
        df = excel_file.parse(sheet, nrows=200)  # cap rows per sheet
        total_rows += len(df)
        parts.append(f"--- Sheet: {sheet} ---")
        if df.empty:
            parts.append("(empty)")
            continue

        # Header row
        parts.append(" | ".join(str(c) for c in df.columns))
        # Data rows
        for _, row in df.head(100).iterrows():
            parts.append(" | ".join(str(v) for v in row.values if v is not None and str(v) != "nan"))

    full_text = "\n".join(parts)

    return {
        "text": full_text,
        "metadata": {
            "filename": filename,
            "file_type": "xlsx",
            "sheet_count": len(sheet_names),
            "sheet_names": sheet_names,
            "rows": total_rows,
        },
    }
